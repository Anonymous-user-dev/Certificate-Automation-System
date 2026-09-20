"""Read and normalize recipient data from Excel workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
import re
from types import MappingProxyType
from typing import Any, Mapping

from openpyxl import load_workbook

from certificate_automation.domain import Recipient


class WorkbookInputError(ValueError):
    """Raised when a workbook cannot safely provide recipient data."""


@dataclass(frozen=True, slots=True)
class WorkbookData:
    """Normalized data read from one worksheet."""

    path: Path
    sheet_name: str
    headers: tuple[str, ...]
    display_headers: Mapping[str, str]
    recipients: tuple[Recipient, ...]

    def __post_init__(self) -> None:
        object.__setattr__(
            self,
            "display_headers",
            MappingProxyType(dict(self.display_headers)),
        )


def normalize_field_name(value: str) -> str:
    """Return a case-insensitive, separator-insensitive field identifier."""

    normalized = re.sub(r"[\W_]+", "_", str(value).strip().casefold())
    return normalized.strip("_")


def list_worksheets(path: Path) -> tuple[str, ...]:
    """List worksheet names without loading cell data into memory."""

    workbook = _open_workbook(path)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def load_workbook_data(path: Path, sheet_name: str) -> WorkbookData:
    """Load and normalize all nonblank recipient rows from one worksheet."""

    workbook = _open_workbook(path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise WorkbookInputError(
                f"Worksheet '{sheet_name}' was not found in '{path.name}'."
            )

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if raw_headers is None or all(_is_blank(value) for value in raw_headers):
            raise WorkbookInputError(
                f"Worksheet '{sheet_name}' does not contain a header row."
            )

        headers, display_headers = _normalize_headers(raw_headers)
        recipients: list[Recipient] = []
        for source_row, row in enumerate(rows, start=2):
            values = tuple(row)
            if all(_is_blank(value) for value in values):
                continue
            normalized_values = {
                header: _normalize_cell(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
            }
            recipients.append(
                Recipient(source_row=source_row, values=normalized_values)
            )

        return WorkbookData(
            path=Path(path),
            sheet_name=sheet_name,
            headers=headers,
            display_headers=display_headers,
            recipients=tuple(recipients),
        )
    finally:
        workbook.close()


def _open_workbook(path: Path):
    path = Path(path)
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise WorkbookInputError(
            f"Excel workbook '{path.name}' could not be read. "
            "It may be damaged or not a valid .xlsx file."
        ) from error


def _normalize_headers(
    raw_headers: tuple[Any, ...],
) -> tuple[tuple[str, ...], Mapping[str, str]]:
    headers: list[str] = []
    display_headers: dict[str, str] = {}

    for column_number, value in enumerate(raw_headers, start=1):
        display_name = "" if value is None else str(value).strip()
        normalized_name = normalize_field_name(display_name)
        if not normalized_name:
            raise WorkbookInputError(
                f"Column {column_number} has a blank column header."
            )
        if normalized_name in display_headers:
            first = display_headers[normalized_name]
            raise WorkbookInputError(
                f"Columns '{first}' and '{display_name}' resolve to the same field name."
            )
        headers.append(normalized_name)
        display_headers[normalized_name] = display_name

    return tuple(headers), display_headers


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
