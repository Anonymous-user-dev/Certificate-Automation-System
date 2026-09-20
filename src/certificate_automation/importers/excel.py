"""Defensive Excel inspection and canonical import."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from hashlib import sha256
from pathlib import Path
from types import MappingProxyType
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from certificate_automation.dataset import (
    Column,
    DataRow,
    SourceSnapshot,
    TabularDataset,
    normalize_column_id,
)


class ExcelImportError(ValueError):
    """An Excel workbook cannot be interpreted without unsafe guessing."""

    def __init__(
        self,
        code: str,
        *,
        path: Path | None = None,
        sheet_name: str | None = None,
        cell: str | None = None,
        ranges: tuple[str, ...] = (),
        parameters: dict[str, str | int] | None = None,
    ) -> None:
        super().__init__(code)
        self.code = code
        self.path = Path(path) if path else None
        self.sheet_name = sheet_name
        self.cell = cell
        self.ranges = tuple(ranges)
        self.parameters = MappingProxyType(dict(parameters or {}))


@dataclass(frozen=True, slots=True)
class ExcelInspection:
    path: Path
    sheet_name: str
    sheet_names: tuple[str, ...]
    hidden_sheets: tuple[str, ...]
    hidden_rows: tuple[int, ...]
    hidden_columns: tuple[str, ...]
    merged_ranges: tuple[str, ...]
    formula_cells: tuple[str, ...]
    formula_cache_missing: tuple[str, ...]

    @property
    def requires_hidden_data_choice(self) -> bool:
        return bool(self.hidden_rows or self.hidden_columns)

    def require_importable(self) -> None:
        if self.merged_ranges:
            raise ExcelImportError(
                "import.excel.merged_data_cells",
                path=self.path,
                sheet_name=self.sheet_name,
                ranges=self.merged_ranges,
            )
        if self.formula_cache_missing:
            raise ExcelImportError(
                "import.excel.formula_cache_missing",
                path=self.path,
                sheet_name=self.sheet_name,
                cell=self.formula_cache_missing[0],
                parameters={"count": len(self.formula_cache_missing)},
            )


def _open_excel(path: Path, *, data_only: bool):
    try:
        return load_workbook(Path(path), read_only=False, data_only=data_only)
    except Exception as error:
        raise ExcelImportError("import.excel.unreadable", path=Path(path)) from error


def _select_sheet(workbook, sheet_name: str | None):
    selected = sheet_name or next(
        (name for name in workbook.sheetnames if workbook[name].sheet_state == "visible"),
        workbook.sheetnames[0] if workbook.sheetnames else None,
    )
    if selected is None or selected not in workbook.sheetnames:
        raise ExcelImportError(
            "import.excel.sheet_missing",
            sheet_name=selected or sheet_name,
        )
    return workbook[selected]


def _used_bounds(worksheet) -> tuple[int, int]:
    max_row = 0
    max_column = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if not _is_blank(cell.value):
                max_row = max(max_row, cell.row)
                max_column = max(max_column, cell.column)
    return max_row, max_column


def _inspect_loaded_sheet(formulas, cached, path: Path) -> ExcelInspection:
    max_row, max_column = _used_bounds(formulas)
    merged = tuple(
        str(cell_range)
        for cell_range in formulas.merged_cells.ranges
        if cell_range.min_row <= max_row
        and cell_range.max_row >= 1
        and cell_range.min_col <= max_column
        and cell_range.max_col >= 1
    )
    formula_cells: list[str] = []
    missing_cache: list[str] = []
    for row in formulas.iter_rows(
        min_row=1,
        max_row=max(max_row, 1),
        min_col=1,
        max_col=max(max_column, 1),
    ):
        for cell in row:
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                formula_cells.append(cell.coordinate)
                if cached[cell.coordinate].value is None:
                    missing_cache.append(cell.coordinate)
    hidden_rows = tuple(
        index
        for index in range(1, max_row + 1)
        if formulas.row_dimensions[index].hidden
    )
    hidden_columns = tuple(
        get_column_letter(index)
        for index in range(1, max_column + 1)
        if formulas.column_dimensions[get_column_letter(index)].hidden
    )
    return ExcelInspection(
        path=Path(path),
        sheet_name=formulas.title,
        sheet_names=tuple(formulas.parent.sheetnames),
        hidden_sheets=tuple(
            sheet.title
            for sheet in formulas.parent.worksheets
            if sheet.sheet_state != "visible"
        ),
        hidden_rows=hidden_rows,
        hidden_columns=hidden_columns,
        merged_ranges=merged,
        formula_cells=tuple(formula_cells),
        formula_cache_missing=tuple(missing_cache),
    )


def inspect_excel(path: Path, sheet_name: str | None = None) -> ExcelInspection:
    """Inspect workbook hazards without importing recipient values."""

    path = Path(path)
    formulas = _open_excel(path, data_only=False)
    cached = _open_excel(path, data_only=True)
    try:
        formula_sheet = _select_sheet(formulas, sheet_name)
        cached_sheet = _select_sheet(cached, formula_sheet.title)
        return _inspect_loaded_sheet(formula_sheet, cached_sheet, path)
    except ExcelImportError as error:
        if error.path is None:
            error.path = path
        raise
    finally:
        formulas.close()
        cached.close()


def import_excel(
    path: Path,
    sheet_name: str,
    *,
    include_hidden: bool = False,
) -> TabularDataset:
    """Import one worksheet only after all unsafe conditions are rejected."""

    path = Path(path)
    formulas = _open_excel(path, data_only=False)
    cached = _open_excel(path, data_only=True)
    try:
        formula_sheet = _select_sheet(formulas, sheet_name)
        cached_sheet = _select_sheet(cached, formula_sheet.title)
        inspection = _inspect_loaded_sheet(formula_sheet, cached_sheet, path)
        inspection.require_importable()
        return _dataset_from_excel_sheets(
            formula_sheet,
            cached_sheet,
            inspection,
            path,
            include_hidden,
        )
    except ExcelImportError as error:
        if error.path is None:
            error.path = path
        if error.sheet_name is None:
            error.sheet_name = sheet_name
        raise
    finally:
        formulas.close()
        cached.close()


def _dataset_from_excel_sheets(
    formulas,
    cached,
    inspection: ExcelInspection,
    path: Path,
    include_hidden: bool,
) -> TabularDataset:
    max_row, max_column = _used_bounds(formulas)
    if max_row == 0 or max_column == 0:
        raise ExcelImportError("import.excel.header_missing", path=path)

    selected_columns = tuple(
        index
        for index in range(1, max_column + 1)
        if include_hidden
        or not formulas.column_dimensions[get_column_letter(index)].hidden
    )
    columns: list[Column] = []
    seen_headers: dict[str, str] = {}
    for index in selected_columns:
        raw_header = cached.cell(1, index).value
        label = "" if raw_header is None else str(raw_header).strip()
        column_id = normalize_column_id(label)
        if not column_id:
            raise ExcelImportError(
                "import.excel.blank_header",
                path=path,
                sheet_name=formulas.title,
                cell=f"{get_column_letter(index)}1",
            )
        if column_id in seen_headers:
            raise ExcelImportError(
                "import.excel.duplicate_header",
                path=path,
                sheet_name=formulas.title,
                cell=f"{get_column_letter(index)}1",
                parameters={"first": seen_headers[column_id], "second": label},
            )
        seen_headers[column_id] = label
        columns.append(Column(column_id, label))

    rows: list[DataRow] = []
    for source_row in range(2, max_row + 1):
        if not include_hidden and formulas.row_dimensions[source_row].hidden:
            continue
        raw_values = [cached.cell(source_row, index).value for index in selected_columns]
        if all(_is_blank(value) for value in raw_values):
            continue
        normalized = {
            column.column_id: _normalize_cell(raw_values[position])
            for position, column in enumerate(columns)
        }
        rows.append(
            DataRow(
                row_id=f"row-{source_row}",
                source_row=source_row,
                values=normalized,
                display_values=normalized,
            )
        )

    source_bytes = path.read_bytes()
    source = SourceSnapshot(
        kind="xlsx",
        label=path.name,
        path=path,
        sha256=sha256(source_bytes).hexdigest(),
        imported_at=datetime.now(timezone.utc),
        options={"sheet_name": formulas.title, "include_hidden": include_hidden},
    )
    return TabularDataset(tuple(columns), tuple(rows), source)


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
