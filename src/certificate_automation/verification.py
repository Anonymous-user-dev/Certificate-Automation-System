"""Independent readability and completeness checks for generated artifacts."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from certificate_automation.template import TemplateInputError, inspect_template


class ArtifactVerificationError(ValueError):
    """Raised when a generated document cannot be trusted as complete."""


def verify_docx(path: Path, forbidden_placeholders: set[str]) -> None:
    """Verify package integrity, Word readability, and replacement completeness."""

    path = Path(path)
    try:
        if not path.is_file() or path.stat().st_size == 0:
            raise ArtifactVerificationError(
                f"Word document '{path.name}' is missing or empty."
            )
        with ZipFile(path, "r") as package:
            damaged = package.testzip()
            if damaged is not None:
                raise ArtifactVerificationError(
                    f"Word document '{path.name}' contains a damaged package member."
                )
        Document(path)
        remaining = set(inspect_template(path).names) & set(forbidden_placeholders)
    except ArtifactVerificationError:
        raise
    except (BadZipFile, OSError, KeyError, ValueError, TemplateInputError) as error:
        raise ArtifactVerificationError(
            f"Word document '{path.name}' could not be opened for verification."
        ) from error

    if remaining:
        joined = ", ".join(f"{{{{{name}}}}}" for name in sorted(remaining))
        raise ArtifactVerificationError(
            f"Word document '{path.name}' still contains an unreplaced placeholder: "
            f"{joined}"
        )


def verify_pdf(path: Path) -> None:
    """Verify PDF signature, parsing, encryption, pages, and page dimensions."""

    path = Path(path)
    try:
        if not path.is_file() or path.stat().st_size == 0:
            raise ArtifactVerificationError(
                f"PDF '{path.name}' is missing or empty and is not a valid PDF."
            )
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise ArtifactVerificationError(
                    f"File '{path.name}' is not a valid PDF."
                )
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ArtifactVerificationError(
                f"PDF '{path.name}' is encrypted and cannot be verified."
            )
        if len(reader.pages) == 0:
            raise ArtifactVerificationError(
                f"PDF '{path.name}' does not contain a page."
            )
        for page_number, page in enumerate(reader.pages, start=1):
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            if width <= 0 or height <= 0:
                raise ArtifactVerificationError(
                    f"PDF '{path.name}' page {page_number} has invalid dimensions."
                )
    except ArtifactVerificationError:
        raise
    except Exception as error:
        raise ArtifactVerificationError(
            f"File '{path.name}' is not a valid PDF and could not be verified."
        ) from error


def verify_pdf_page_count(path: Path, expected_pages: int) -> int:
    """Reopen a PDF and require the exact expected page count."""

    verify_pdf(path)
    try:
        actual_pages = len(PdfReader(Path(path), strict=True).pages)
    except Exception as error:
        raise ArtifactVerificationError(
            f"PDF '{Path(path).name}' page count could not be verified."
        ) from error
    if actual_pages != expected_pages:
        raise ArtifactVerificationError(
            f"PDF '{Path(path).name}' has {actual_pages} pages; expected {expected_pages}."
        )
    return actual_pages


def pdf_page_count(path: Path) -> int:
    """Return the page count only after the PDF passes strict verification."""

    verify_pdf(path)
    try:
        return len(PdfReader(Path(path), strict=True).pages)
    except Exception as error:
        raise ArtifactVerificationError(
            f"PDF '{Path(path).name}' page count could not be read."
        ) from error
