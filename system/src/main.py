from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent

DATA_FILE = PROJECT_ROOT / "data" / "students.xlsx"

def load_students(file_path: Path):
    workbook = load_workbook(file_path, data_only=True)

    worksheet = workbook.active

    students = []

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        full_name, award, date = row

        students.append(
            {
                "row_number": row_number,
                "full_name": full_name,
                "award": award,
                "date": date,
            }
        )

    return students


students = load_students(DATA_FILE)

for student in students:
    print(student)