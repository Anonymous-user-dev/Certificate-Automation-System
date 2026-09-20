from __future__ import annotations

from zipfile import ZipFile

import pytest
from docx import Document

from certificate_automation.template import (
    TemplateInputError,
    TemplateRenderError,
    inspect_template,
    render_template,
)
from fixtures import docx_factory, ooxml_docx_factory


def document_text(path):
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(part for part in parts if part)


def test_replaces_placeholder_split_across_runs(docx_factory, tmp_path):
    source = docx_factory(
        paragraph_runs=[["Issued to ", "{{FULL_", "NAME}}", "."]]
    )
    output = tmp_path / "out.docx"

    render_template(source, output, {"FULL_NAME": "Ana García"})

    assert document_text(output) == "Issued to Ana García."


def test_discovers_placeholders_in_text_boxes(ooxml_docx_factory):
    source = ooxml_docx_factory(text_box="{{CERTIFICATE_ID}}")

    inspection = inspect_template(source)

    assert inspection.names == ("CERTIFICATE_ID",)
    assert inspection.placeholders[0].occurrences == 1


def test_replaces_text_box_placeholder(ooxml_docx_factory, tmp_path):
    source = ooxml_docx_factory(text_box="ID {{CERTIFICATE_ID}}")
    output = tmp_path / "text-box.docx"

    render_template(source, output, {"CERTIFICATE_ID": "CERT-2026-001"})

    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "ID CERT-2026-001" in document_xml
    assert inspect_template(output).names == ()


def test_replaces_repeated_placeholders_in_body_table_header_and_footer(
    docx_factory,
    tmp_path,
):
    source = docx_factory(
        paragraph_runs=[["{{FULL_NAME}} / {{FULL_NAME}}"]],
        table_text="Award: {{AWARD}}",
        header_text="Recipient {{FULL_NAME}}",
        footer_text="Date {{DATE}}",
    )
    output = tmp_path / "all-parts.docx"

    render_template(
        source,
        output,
        {
            "FULL_NAME": "Ana García",
            "AWARD": "Gold",
            "DATE": "2026-09-20",
        },
    )

    text = document_text(output)
    assert "Ana García / Ana García" in text
    assert "Award: Gold" in text
    assert "Recipient Ana García" in text
    assert "Date 2026-09-20" in text
    assert inspect_template(output).names == ()


def test_inspection_reports_first_seen_order_and_occurrence_count(docx_factory):
    source = docx_factory(
        paragraph_runs=[
            ["{{AWARD}} then {{FULL_NAME}}"],
            ["Again {{AWARD}}"],
        ]
    )

    inspection = inspect_template(source)

    assert inspection.names == ("AWARD", "FULL_NAME")
    assert [item.occurrences for item in inspection.placeholders] == [2, 1]


@pytest.mark.parametrize(
    "text",
    ["{{}}", "{{   }}", "Start {{FULL_NAME", "End FULL_NAME}}", "{{A{{B}}"],
)
def test_malformed_placeholder_is_rejected(docx_factory, text):
    source = docx_factory(paragraph_runs=[[text]])

    with pytest.raises(TemplateInputError, match="malformed placeholder"):
        inspect_template(source)


def test_unknown_replacement_key_is_rejected(docx_factory, tmp_path):
    source = docx_factory(paragraph_runs=[["{{FULL_NAME}}"]])

    with pytest.raises(TemplateRenderError, match="not present"):
        render_template(source, tmp_path / "out.docx", {"AWARD": "Gold"})


def test_source_template_is_not_modified(docx_factory, tmp_path):
    source = docx_factory(paragraph_runs=[["{{FULL_NAME}}"]])
    before = source.read_bytes()

    render_template(source, tmp_path / "out.docx", {"FULL_NAME": "Ana"})

    assert source.read_bytes() == before


def test_rendered_package_reopens_and_has_valid_zip_members(docx_factory, tmp_path):
    source = docx_factory(paragraph_runs=[["Certificate for {{FULL_NAME}}"]])
    output = tmp_path / "verified.docx"

    render_template(source, output, {"FULL_NAME": "Ana"})

    assert Document(output).paragraphs[-1].text == "Certificate for Ana"
    with ZipFile(output) as package:
        assert package.testzip() is None


def test_invalid_docx_has_plain_language_error(tmp_path):
    source = tmp_path / "invalid.docx"
    source.write_bytes(b"not a Word document")

    with pytest.raises(TemplateInputError, match="could not be read"):
        inspect_template(source)
