from __future__ import annotations

import platform
import os
from datetime import datetime, timezone
from pathlib import Path
import subprocess
import sys

import pytest
from docx import Document

from certificate_automation.dataset import Column, DataRow, SourceSnapshot, TabularDataset
from certificate_automation.mapping import ColumnValue, MappingPlan
from certificate_automation.template import inspect_template
from certificate_automation.template_health import TemplateHealthService
from certificate_automation.verification import verify_pdf
from certificate_automation.word import WordPdfConverter


pytestmark = [
    pytest.mark.word_integration,
    pytest.mark.skipif(platform.system() != "Windows", reason="requires Windows"),
]


def test_installed_word_produces_readable_one_page_pdf(tmp_path):
    source = tmp_path / "certificate.docx"
    destination = tmp_path / "certificate.pdf"
    document = Document()
    document.add_heading("Certificate of Achievement", level=1)
    document.add_paragraph("Presented to Ana García")
    document.save(source)

    converter = WordPdfConverter(max_attempts=2)
    availability = converter.is_available()
    assert availability.available, availability.message

    converter.convert(source, destination)

    verify_pdf(destination)


def test_word_representative_layout_review_uses_new_verified_pdf(tmp_path):
    source = tmp_path / "layout-template.docx"
    document = Document()
    document.add_paragraph("Certificate awarded to {{FULL_NAME}}")
    document.save(source)
    template = inspect_template(source)
    dataset = TabularDataset(
        (Column("name", "Name"),),
        (
            DataRow("first", 2, {"name": "Ana García"}),
            DataRow("longest", 3, {"name": "Анастасия Петровна Кузнецова"}),
            DataRow("last", 4, {"name": "李明"}),
        ),
        SourceSnapshot("manual", "Manual", None, "a" * 64, datetime.now(timezone.utc)),
    )
    converter = WordPdfConverter(max_attempts=2)
    availability = converter.is_available()
    assert availability.available, availability.message

    result = TemplateHealthService(converter, tmp_path / "previews").render_representatives(
        dataset, template, MappingPlan({"FULL_NAME": ColumnValue("name")}), expected_pages=1
    )

    assert result.ready, result.issues
    assert {item.row_id for item in result.previews} == {"first", "longest", "last"}
    assert all(item.page_count == 1 and item.pdf_path.is_file() for item in result.previews)


def test_word_com_teardown_has_no_fatal_rpc_output(tmp_path):
    source = tmp_path / "teardown.docx"
    destination = tmp_path / "teardown.pdf"
    document = Document()
    document.add_paragraph("Teardown verification")
    document.save(source)
    script = """
from pathlib import Path
import sys
from certificate_automation.word import WordPdfConverter
from certificate_automation.verification import verify_pdf
source, destination = map(Path, sys.argv[1:])
WordPdfConverter(max_attempts=1).convert(source, destination)
verify_pdf(destination)
"""
    environment = os.environ.copy()
    environment["PYTHONFAULTHANDLER"] = "1"
    source_root = str(Path(__file__).resolve().parents[2] / "src")
    environment["PYTHONPATH"] = os.pathsep.join(
        part
        for part in (source_root, environment.get("PYTHONPATH", ""))
        if part
    )

    stdout_path = tmp_path / "subprocess.stdout.txt"
    stderr_path = tmp_path / "subprocess.stderr.txt"
    result = subprocess.run(
        [sys.executable, "-c", script, str(source), str(destination)],
        capture_output=True,
        text=True,
        env=environment,
        timeout=60,
        check=False,
    )
    stdout_path.write_text(result.stdout, encoding="utf-8")
    stderr_path.write_text(result.stderr, encoding="utf-8")
    error_output = result.stderr

    assert result.returncode == 0, error_output
    assert "fatal exception" not in error_output.casefold(), error_output
